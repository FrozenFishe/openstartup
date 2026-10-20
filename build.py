#!/usr/bin/env python3
"""Build OpenStartup from the research files.

research/countries/*.json  (one array per regional batch, see research/SCHEMA.md)
research/global/accelerators.json, research/global/resources.json
research/wikipedia/vcs.json  (Wikipedia's List of venture capital firms, extra investors per country)
research/wikipedia/pe.json   (Wikipedia's List of private equity firms, a separate "pe" list per country)
        |
        v
data/*.csv, data/openstartup.json, data/openstartup.xlsx
webapp/atlas.html  (page fragment, also published as an Artifact)
docs/index.html    (GitHub Pages, full document) + docs/data/ copies

Needs Python 3; openpyxl for the workbook (skipped if missing).
Map geometry, ALIAS and POINTS are taken from OpenCharity's build_atlas.py.
"""
import csv
import datetime as dt
import json
import re
import shutil
import sys
from pathlib import Path

BASE = Path(__file__).parent
RESEARCH = BASE / "research"
DATA = BASE / "data"
GEO = BASE / "webapp" / "geo" / "countries-50m.json"
TEMPLATE = BASE / "atlas_template.html"
OUT = BASE / "webapp" / "atlas.html"
DOCS = BASE / "docs"
TODAY = dt.date(2026, 10, 20).isoformat()  # the date deadlines were checked against

REGION_OF_BATCH = {
    "b01_usa": "North America", "b02_canada_uk_ie": None, "b03_latam_big": "Latin America & Caribbean",
    "b04_latam_carib": "Latin America & Caribbean", "b05_weu": "Europe", "b06_nordic_baltic_micro": "Europe",
    "b07_ceu_east": "Europe", "b08_mena": "Middle East & North Africa", "b09_south_asia": "South Asia",
    "b10_east_asia": "East Asia", "b11_sea": "Southeast Asia", "b12_oceania": "Oceania",
    "b13_africa_east_south": "Sub-Saharan Africa", "b14_africa_west_central": "Sub-Saharan Africa",
    "b15_central_asia_caucasus": "Central Asia & Caucasus",
}
REGION_FIX = {"Canada": "North America", "United Kingdom": "Europe", "Ireland": "Europe",
              "Sudan": "Middle East & North Africa", "Greenland": "North America"}

# dataset name -> Natural Earth feature name (from OpenCharity build_atlas.py)
ALIAS = {
    "United States": "United States of America", "Trinidad & Tobago": "Trinidad and Tobago",
    "Antigua & Barbuda": "Antigua and Barb.", "St. Kitts & Nevis": "St. Kitts and Nevis", "St. Lucia": "Saint Lucia",
    "St. Vincent & Grenadines": "St. Vin. and Gren.", "Dominican Republic": "Dominican Rep.", "Macau": "Macao",
    "Solomon Islands": "Solomon Is.", "Marshall Islands": "Marshall Is.", "Central African Republic": "Central African Rep.",
    "Republic of the Congo": "Congo", "Congo (DRC)": "Dem. Rep. Congo", "Equatorial Guinea": "Eq. Guinea",
    "Sao Tome & Principe": "São Tomé and Principe", "Eswatini": "eSwatini", "Bosnia & Herzegovina": "Bosnia and Herz.",
    "North Macedonia": "Macedonia", "Vatican City": "Vatican", "Western Sahara": "W. Sahara", "South Sudan": "S. Sudan",
}
POINTS = {"Tuvalu": [179.2, -8.5]}  # no 1:50m polygon: [lon, lat]

TIERS = {"global_hub", "established", "emerging", "nascent"}
LISTS = ("grants", "accelerators", "vcs", "pe", "resources")
FIELDS = {
    "grants": ["name", "agency", "amount", "amount_usd_max", "eligibility", "deadline", "status", "url"],
    "accelerators": ["name", "city", "type", "investment", "sectors", "deadline", "next_cohort", "status", "url"],
    "vcs": ["name", "city", "stage", "check_size", "aum", "sectors", "type", "url", "source"],
    "pe": ["name", "city", "aum", "founded", "url", "source"],
    "resources": ["name", "type", "url", "note"],
}
G_ACC = ["name", "organizer", "type", "locations", "countries_eligible", "investment", "equity", "sectors",
         "deadline", "deadline_note", "next_cohort", "status", "url"]
G_RES = ["name", "category", "coverage", "free", "url", "note"]
URL_RE = re.compile(r"^https?://[^\s/$.?#].[^\s]*$", re.I)


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        return v
    return re.sub(r"\s+", " ", str(v)).strip()


def tidy(item, fields, problems, where):
    out = {f: clean(item.get(f)) for f in fields}
    if not out.get("name"):
        problems.append(f"{where}: entry without a name dropped")
        return None
    if out.get("url") and not URL_RE.match(out["url"]):
        problems.append(f"{where}: bad url for {out['name']!r}: {out['url']!r}")
        out["url"] = ""
    return out


def dedupe(items):
    seen, out = set(), []
    for x in items:
        key = (x["name"].lower(), (x.get("url") or "").lower().rstrip("/"))
        if key in seen:
            continue
        seen.add(key)
        out.append(x)
    return out


def load_wikipedia(kind):
    """Country -> firms from a Wikipedia list (research/wikipedia/<kind>.json), shaped like that list's entries."""
    f = RESEARCH / "wikipedia" / f"{kind}.json"
    if not f.exists():
        return {}
    src = json.loads(f.read_text(encoding="utf-8"))
    label = {"vcs": "Wikipedia: List of venture capital firms", "pe": "Wikipedia: List of private equity firms"}[kind]
    out = {}
    for x in src["firms"]:
        row = {"name": x["name"], "city": x.get("city", ""), "aum": x.get("aum", ""), "founded": x.get("founded", ""),
               "url": x["url"], "source": label}
        if kind == "vcs":
            row.update(stage="", check_size="unknown", sectors="", type="vc")
        out.setdefault(x["country"], []).append(row)
    return out


def load_countries(problems):
    wiki = {k: load_wikipedia(k) for k in ("vcs", "pe")}
    batches = json.loads((RESEARCH / "batches.json").read_text(encoding="utf-8"))
    want = {c: b for b, cs in batches.items() for c in cs}
    got = {}
    for f in sorted((RESEARCH / "countries").glob("*.json")):
        try:
            rows = json.loads(f.read_text(encoding="utf-8"))
        except json.JSONDecodeError as e:
            problems.append(f"{f.name}: invalid JSON ({e}); skipped")
            continue
        for r in rows if isinstance(rows, list) else []:
            name = clean(r.get("country"))
            if name not in want:
                problems.append(f"{f.name}: unknown country {name!r}")
                continue
            got[name] = r
    out = []
    for name, batch in want.items():
        r = got.get(name)
        region = REGION_FIX.get(name) or REGION_OF_BATCH[batch]
        if r is None:
            problems.append(f"{name}: not researched yet (stub row)")
            r = {"country": name, "confidence": "", "summary": ""}
        tier = clean(r.get("ecosystem_tier")).lower().replace(" ", "_")
        d = {
            "country": name, "iso2": clean(r.get("iso2")), "region": region,
            "tier": tier if tier in TIERS else "nascent",
            "summary": clean(r.get("summary")), "startup_visa": clean(r.get("startup_visa")),
            "company_setup": clean(r.get("company_setup")), "notes": clean(r.get("notes")),
            "confidence": clean(r.get("confidence")).lower(),
            "sources": [u for u in (clean(s) for s in r.get("sources") or []) if URL_RE.match(u)],
            "researched": name in got,
        }
        for k in LISTS:
            items = [tidy(x, FIELDS[k], problems, f"{name}/{k}") for x in (r.get(k) or []) if isinstance(x, dict)]
            if k in wiki:  # research entries first, then Wikipedia firms not already listed
                items += [tidy(x, FIELDS[k], problems, f"{name}/wikipedia-{k}") for x in wiki[k].get(name, [])]
            d[k] = dedupe([x for x in items if x])
        out.append(d)
    return sorted(out, key=lambda d: d["country"])


def load_global(problems):
    res = {}
    for key, fields in (("accelerators", G_ACC), ("resources", G_RES)):
        f = RESEARCH / "global" / f"{key}.json"
        rows = []
        if f.exists():
            try:
                rows = json.loads(f.read_text(encoding="utf-8"))
            except json.JSONDecodeError as e:
                problems.append(f"global/{key}.json: invalid JSON ({e})")
        else:
            problems.append(f"global/{key}.json: missing")
        items = [tidy(x, fields, problems, f"global/{key}") for x in rows if isinstance(x, dict)]
        res[key] = dedupe([x for x in items if x])
    return res


def write_csv(path, header, rows):
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        w.writerows(rows)


def write_data(countries, glob):
    DATA.mkdir(exist_ok=True)
    files = []
    head = ["country", "iso2", "region", "ecosystem_tier", "accelerators", "vcs", "private_equity", "grants", "resources",
            "startup_visa", "company_setup", "summary", "notes", "confidence", "sources"]
    write_csv(DATA / "countries.csv", head, [[d["country"], d["iso2"], d["region"], d["tier"],
              *(len(d[k]) for k in ("accelerators", "vcs", "pe", "grants", "resources")),
              d["startup_visa"], d["company_setup"], d["summary"], d["notes"], d["confidence"], "; ".join(d["sources"])]
              for d in countries])
    files.append("countries.csv")
    for k in LISTS:
        write_csv(DATA / f"{k}.csv", ["country", "region", *FIELDS[k]],
                  [[d["country"], d["region"], *(x[f] for f in FIELDS[k])] for d in countries for x in d[k]])
        files.append(f"{k}.csv")
    write_csv(DATA / "global_programs.csv", G_ACC, [[x[f] for f in G_ACC] for x in glob["accelerators"]])
    write_csv(DATA / "global_resources.csv", G_RES, [[x[f] for f in G_RES] for x in glob["resources"]])
    files += ["global_programs.csv", "global_resources.csv"]
    (DATA / "openstartup.json").write_text(json.dumps({"built": TODAY, "countries": countries, "global": glob},
                                                      ensure_ascii=False, indent=1), encoding="utf-8")
    files.append("openstartup.json")
    if write_xlsx(countries, glob):
        files.append("openstartup.xlsx")
    return files


def write_xlsx(countries, glob):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed: skipping openstartup.xlsx")
        return False
    head_fill, head_font = PatternFill("solid", start_color="1D4ED8"), Font(color="FFFFFF", bold=True)
    tier_fill = {"global_hub": "0B2A6F", "established": "1D4ED8", "emerging": "93B8FD", "nascent": "DBE7FF"}
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title, header, rows, widths=None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 18)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        return ws

    ws = sheet("Countries", ["Country", "Region", "Ecosystem", "Accelerators", "VCs", "Private equity", "Grants", "Resources",
                             "Startup visa", "Company setup", "Summary", "Confidence"],
               [[d["country"], d["region"], d["tier"], len(d["accelerators"]), len(d["vcs"]), len(d["pe"]), len(d["grants"]),
                 len(d["resources"]), d["startup_visa"], d["company_setup"], d["summary"], d["confidence"]] for d in countries],
               {"Startup visa": 40, "Company setup": 40, "Summary": 60})
    for row in ws.iter_rows(min_row=2):
        col = tier_fill.get(row[2].value)
        if col:
            row[2].fill = PatternFill("solid", start_color=col)
            if col in ("0B2A6F", "1D4ED8"):
                row[2].font = Font(color="FFFFFF")
    wide = {"name": 34, "eligibility": 40, "amount": 26, "investment": 28, "url": 40, "note": 40, "sectors": 26}
    for k in LISTS:
        sheet({"vcs": "VCs", "pe": "Private equity"}.get(k, k.capitalize()), ["country", *FIELDS[k]],
              [[d["country"], *(x[f] for f in FIELDS[k])] for d in countries for x in d[k]], wide)
    sheet("Global programs", G_ACC, [[x[f] for f in G_ACC] for x in glob["accelerators"]], wide)
    sheet("Databases", G_RES, [[x[f] for f in G_RES] for x in glob["resources"]], wide)
    wb.save(DATA / "openstartup.xlsx")
    return True


def build_page(countries, glob, files):
    geo = json.loads(GEO.read_text(encoding="utf-8"))
    geo["objects"]["countries"]["geometries"] = [g for g in geo["objects"]["countries"]["geometries"]
                                                 if g["properties"]["name"] != "Antarctica"]
    names = {g["properties"]["name"] for g in geo["objects"]["countries"]["geometries"]}
    missing = [d["country"] for d in countries if ALIAS.get(d["country"], d["country"]) not in names and d["country"] not in POINTS]
    if missing:
        raise SystemExit(f"no map shape for: {missing}")

    def js(obj):
        return json.dumps(obj, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")

    slim = [{k: v for k, v in d.items() if k != "researched"} for d in countries]
    html = (TEMPLATE.read_text(encoding="utf-8")
            .replace("__DATA__", js(slim)).replace("__GLOBAL__", js(glob)).replace("__GEO__", js(geo))
            .replace("__ALIAS__", js(ALIAS)).replace("__POINTS__", js(POINTS))
            .replace("__META__", js({"today": TODAY, "files": files})))
    OUT.write_text(html, encoding="utf-8")
    # GitHub Pages: the same page in a full document (the Artifact host adds this skeleton itself)
    head, body = html.split('<div class="wrap">', 1)
    doc = ('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
           '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
           '<style>[hidden]{display:none!important} body{margin:0} img{max-width:100%}</style>\n'
           f'{head}</head>\n<body>\n<div class="wrap">{body}</body>\n</html>\n')
    DOCS.mkdir(exist_ok=True)
    (DOCS / "index.html").write_text(doc, encoding="utf-8")
    (DOCS / ".nojekyll").write_text("", encoding="utf-8")
    shutil.copytree(DATA, DOCS / "data", dirs_exist_ok=True)


def main():
    problems = []
    countries = load_countries(problems)
    glob = load_global(problems)
    files = write_data(countries, glob)
    build_page(countries, glob, files)
    done = sum(d["researched"] for d in countries)
    tot = {k: sum(len(d[k]) for d in countries) for k in LISTS}
    print(f"{len(countries)} countries ({done} researched) · " + " · ".join(f"{v} {k}" for k, v in tot.items())
          + f" · {len(glob['accelerators'])} global programs · {len(glob['resources'])} databases")
    print(f"wrote {OUT.relative_to(BASE)}, docs/index.html and data/: {', '.join(files)}")
    (RESEARCH / "build_problems.txt").write_text("\n".join(problems) + "\n", encoding="utf-8")
    if problems:
        print(f"{len(problems)} problems listed in research/build_problems.txt")
    return 0


if __name__ == "__main__":
    sys.exit(main())
